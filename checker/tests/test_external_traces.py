"""내보내도 되는 파일인지 보는 규칙과, 흔적을 걷어내는 도구.

다른 규칙 종류와 묻는 것이 다르다. 나머지는 "이 회사 템플릿과 룰대로 썼나" 를 보지만
이것은 "이거 내보내도 되나" 를 본다. 서로 독립이라 템플릿을 완벽히 따른 문서에도 흔적이
있을 수 있다.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from checker.cleaner import clean, default_output
from checker.ooxml import document_properties, is_ooxml, named_styles
from checker.rules import get


def rule(path: Path, params: dict | None = None):
    return get("external_traces")(path, params or {})


@pytest.fixture()
def traces(sample: Path) -> Path:
    return sample / "docs" / "흔적"


def test_깨끗한_파일은_통과한다(traces):
    assert rule(traces / "깨끗한.xlsx") == []


def test_흔적_다섯_가지를_모두_잡는다(traces):
    found = {v.message for v in rule(traces / "흔적있는.xlsx")}
    joined = " ".join(found)
    assert "명명 스타일" in joined       # 복붙으로 쌓인 스타일
    assert "customXml" in joined        # SharePoint 가 심은 자리
    assert "dc:creator" in joined       # 작성자
    assert "cp:lastModifiedBy" in joined  # 최종 수정자
    assert "주석 작성자" in joined


def test_어디에_무엇이_있는지_짚는다(traces):
    """파일 안쪽 경로까지 나와야 사람이 조치할 수 있다."""
    violations = rule(traces / "흔적있는.xlsx")
    creator = next(v for v in violations if "dc:creator" in v.message)
    assert creator.actual == "가짜작성자"


def test_이름공간_속성이_붙어도_잡는다(traces):
    """같은 요소라도 만든 도구에 따라 `<dc:creator xmlns:dc=...>` 처럼 속성이 붙는다.

    속성을 허용하지 않으면 도구에 따라 작성자 정보를 통째로 놓친다. 실제로 그 버그가
    있었고, openpyxl 이 만든 파일이 전부 빠져나갔다.
    """
    with zipfile.ZipFile(traces / "흔적있는.xlsx") as zf:
        raw = zf.read("docProps/core.xml").decode("utf-8")
        assert 'xmlns:dc=' in raw, "픽스처가 속성 붙은 형태여야 이 검사가 의미 있다"
        assert document_properties(zf).get("dc:creator") == "가짜작성자"


def test_임계를_올리면_스타일은_넘어간다(traces):
    found = {v.message for v in rule(traces / "흔적있는.xlsx", {"최대_명명_스타일": 1000})}
    assert not any("명명 스타일" in m for m in found)


def test_금지어를_따로_지정할_수_있다(traces):
    """구조적 신호로 잡히지 않는 것을 고객사별로 더하는 용도다."""
    found = {v.message for v in rule(traces / "흔적있는.xlsx", {"금지어": ["가짜거래처"]})}
    assert any("가짜거래처" in m for m in found)


def test_오피스_파일이_아니면_관여하지_않는다(sample):
    assert rule(sample / "templates" / "제안서.md") == []


# --- clean ----------------------------------------------------------------

def test_걷어내면_검사를_통과한다(traces, tmp_path):
    """검사만 하고 지우는 수단이 없으면 조치할 수 없는 경고만 남는다.

    명명 스타일 수만 개를 엑셀 UI 에서 지울 방법이 없다.
    """
    out = tmp_path / "정리본.xlsx"
    result = clean(traces / "흔적있는.xlsx", out)
    assert result.removed
    assert rule(out) == []


def test_걷어내도_시트와_서식이_남는다(traces, tmp_path):
    """openpyxl 로 열어 저장하면 이미지가 사라진다. 그래서 zip 항목을 직접 다룬다."""
    import openpyxl

    out = tmp_path / "정리본.xlsx"
    clean(traces / "흔적있는.xlsx", out)

    before = openpyxl.load_workbook(traces / "흔적있는.xlsx")
    after = openpyxl.load_workbook(out)
    assert after.sheetnames == before.sheetnames
    assert after.active["A1"].value == before.active["A1"].value

    with zipfile.ZipFile(traces / "흔적있는.xlsx") as a, zipfile.ZipFile(out) as b:
        assert len(b.namelist()) == len(a.namelist())  # 항목을 지우지 않는다
        # 기본 스타일은 남는다
        assert named_styles(b)[0] == named_styles(a)[0]


def test_원본을_덮지_않는다(traces, tmp_path):
    src = tmp_path / "원본.xlsx"
    src.write_bytes((traces / "흔적있는.xlsx").read_bytes())
    before = src.read_bytes()
    clean(src, default_output(src))
    assert src.read_bytes() == before
    assert default_output(src).exists()


def test_무엇을_지웠는지_보고한다(traces, tmp_path):
    result = clean(traces / "흔적있는.xlsx", tmp_path / "정리본.xlsx")
    joined = result.summary()
    assert "명명 스타일" in joined
    assert "가짜작성자" in joined


def test_오피스_파일이_아니면_거부한다(sample, tmp_path):
    with pytest.raises(ValueError):
        clean(sample / "templates" / "제안서.md", tmp_path / "x.md")


def test_엔진은_이_규칙도_모른다():
    """일곱 번째 규칙 종류를 붙이면서 engine.py 가 바뀌지 않았다."""
    import inspect

    from checker import engine

    assert "external_traces" not in inspect.getsource(engine)


# --- 템플릿 없는 유형 ------------------------------------------------------

def _rules_dir(tmp_path: Path, body: str) -> Path:
    d = tmp_path / "회사" / "rules"
    d.mkdir(parents=True)
    (d / "규칙.yaml").write_text(body, encoding="utf-8")
    return d


def test_템플릿을_읽지_않는_규칙만_있으면_템플릿이_없어도_된다(tmp_path):
    """템플릿 자체를 검사하려면 이 구분이 필요하다. 템플릿의 템플릿은 없다.

    오염된 템플릿이 모든 산출물을 오염시키므로 거기가 발원지다.
    """
    from checker.loader import load

    rules = _rules_dir(tmp_path, "관할: 'templates/**'\n규칙:\n  - 종류: external_traces\n")
    types = load(rules)
    assert types[0].template == ""
    assert [r.kind for r in types[0].rules] == ["external_traces"]


def test_템플릿을_읽는_규칙이_있으면_템플릿을_요구한다(tmp_path):
    from checker.loader import load
    from checker.model import ConfigError

    rules = _rules_dir(tmp_path, "관할: 'docs/**'\n규칙:\n  - 종류: required_sections\n")
    with pytest.raises(ConfigError, match="템플릿"):
        load(rules)


def test_템플릿_없는_유형의_리포트는_템플릿이_null_이다(tmp_path):
    from checker.engine import check
    from checker.loader import load

    rules = _rules_dir(tmp_path, "관할: 'templates/**'\n규칙:\n  - 종류: external_traces\n")
    doc = tmp_path / "회사" / "templates" / "무언가.md"
    doc.parent.mkdir(parents=True)
    doc.write_text("# 제목\n", encoding="utf-8")

    report = check([(doc, "templates/무언가.md")], load(rules))
    assert report["files"][0]["template"] is None
    assert report["files"][0]["status"] == "pass"
