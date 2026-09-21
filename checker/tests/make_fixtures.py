"""픽스처를 만든다. 결과물은 저장소에 커밋하지 않는다.

docx·xlsx·pptx 는 바이너리라 커밋하면 리뷰에서 diff 를 볼 수 없고, "이 픽스처가 정말
의도한 위반을 담고 있나" 를 아무도 확인할 수 없다. 생성 스크립트로 두면 무엇이 들어 있는지가
코드로 읽히고 고칠 때 diff 가 남는다.

만들어지는 구조는 실제 문서 저장소와 같다. 회사 폴더 하나가 templates/ rules/ docs/ 를
들고 있고, 검사는 그 바깥에서 부른다.
"""
from __future__ import annotations

import shutil
from pathlib import Path

ROOT = Path(__file__).parent / "fixtures" / "client-docs" / "sample"

# 유형마다 통과 1건, 의도적 위반 1건, 경계 1건을 짝으로 둔다. 통과만 넣으면 검사기가
# 아무것도 하지 않아도 테스트가 초록불이 된다.

PROPOSAL_TEMPLATE = r"""# 제안서

## 개요

## 범위

## 기능 요구사항

## 일정

## 부록
"""

PROPOSAL_PASS = r"""# 로그인 개선 제안서

## 개요
로그인 절차를 줄인다.

## 범위
웹과 모바일.

## 기능 요구사항
- 소셜 로그인

## 일정
2026 4분기.

## 부록
없음.
"""

# '일정' 이 빠졌다. 템플릿에 있으므로 위반이어야 한다.
PROPOSAL_MISSING = r"""# 결제 개선 제안서

## 개요
결제를 고친다.

## 범위
웹.

## 기능 요구사항
- 간편결제
"""

# '부록' 만 빠졌다. 선택 섹션으로 지정했으므로 통과해야 한다. 경계 사례다.
PROPOSAL_BOUNDARY = r"""# 배송 개선 제안서

## 개요
배송을 고친다.

## 범위
전국.

## 기능 요구사항
- 당일배송

## 일정
2027 1분기.
"""

RULES_PROPOSAL = r"""템플릿: ../templates/제안서.md
관할: "docs/제안서/**"

규칙:
  - 종류: filename
    패턴: '^\d{8}_.+_제안서\.md$'
  - 종류: location
    있어야_할_곳: docs/제안서
  - 종류: required_sections
    선택_섹션: [부록]
"""

RULES_SPEC = r"""템플릿: ../templates/개발사양서.xlsx
관할: "docs/개발사양서/**"

규칙:
  - 종류: required_sheets
    필수: [Title, 개발 기준, 테이블 설계]
  - 종류: table_headers
    시트: 테이블 설계
    헤더_행: 12
"""

RULES_DECK = r"""템플릿: ../templates/발표자료.pptx
관할: "docs/발표자료/**"

규칙:
  - 종류: slide_layouts
"""


def _write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    # 인코딩을 명시한다. Windows 기본 코드페이지에 맡기면 한글이 깨진다.
    path.write_text(text, encoding="utf-8")


def build_md() -> None:
    _write(ROOT / "templates" / "제안서.md", PROPOSAL_TEMPLATE)
    _write(ROOT / "rules" / "제안서.yaml", RULES_PROPOSAL)
    d = ROOT / "docs" / "제안서"
    _write(d / "20260921_로그인_제안서.md", PROPOSAL_PASS)
    _write(d / "20260921_결제_제안서.md", PROPOSAL_MISSING)
    _write(d / "20260921_배송_제안서.md", PROPOSAL_BOUNDARY)
    _write(d / "제안서최종.md", PROPOSAL_PASS)  # 내용은 맞고 이름만 틀렸다
    _write(ROOT / "README.md", "검사 대상이 아닌 파일. 조용히 무시되어야 한다.\n")


def _spec_sheet(ws, headers, rows):
    # 실제 템플릿처럼 헤더를 12행에 둔다. 위쪽은 레이블-값 쌍이라 자동 탐지가 헛짚는 자리다.
    ws["A5"], ws["B5"] = "테이블ID", "ZSAMPLE01"
    ws["A6"], ws["B6"] = "설명", "예시 테이블"
    for i, h in enumerate(headers, start=1):
        ws.cell(12, i, h)
    for r, row in enumerate(rows, start=13):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)


def build_xlsx() -> None:
    import openpyxl

    _write(ROOT / "rules" / "개발사양서.yaml", RULES_SPEC)

    headers = ["순번", "칼럼ID", "칼럼명", "PK", "데이터타입", "길이"]

    def book(sheets, hdr):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name in sheets:
            ws = wb.create_sheet(name)
            if name == "테이블 설계":
                _spec_sheet(ws, hdr, [[1, "ID", "식별자", "O", "CHAR", 10]])
            else:
                ws["A1"] = name
        return wb

    full = ["Title", "개발 기준", "테이블 설계"]
    book(full, headers).save(ROOT / "templates" / "개발사양서.xlsx")

    d = ROOT / "docs" / "개발사양서"
    d.mkdir(parents=True, exist_ok=True)
    # 통과: 시트도 헤더도 템플릿과 같다
    book(full, headers).save(d / "통과.xlsx")
    # 위반: '테이블 설계' 시트가 없다
    book(["Title", "개발 기준"], headers).save(d / "시트없음.xlsx")
    # 위반: 헤더 한 칸이 다르다
    book(full, ["순번", "컬럼ID", "칼럼명", "PK", "데이터타입", "길이"]).save(d / "헤더다름.xlsx")
    # 경계: 프로그램마다 붙는 시트가 하나 더 있다. 필수 시트는 다 있으므로 통과해야 한다
    book(full + ["화면 레이아웃"], headers).save(d / "시트추가.xlsx")


def build_pptx() -> None:
    from pptx import Presentation

    _write(ROOT / "rules" / "발표자료.yaml", RULES_DECK)

    def deck(path, layout_indexes):
        prs = Presentation()
        for i in layout_indexes:
            prs.slides.add_slide(prs.slide_layouts[i])
        prs.save(str(path))

    # 템플릿은 레이아웃을 셋만 정의하게 깎는다. 실제 공식 템플릿이 그런 모양이었다
    # (DEFAULT / 제목 슬라이드 / 1_Picture with Caption 셋뿐). 깎지 않으면 파워포인트
    # 기본 마스터가 11개를 다 정의해서 무엇을 써도 허용 범위에 들어가 버린다.
    keep = (0, 1, 5)
    prs = Presentation()
    for i in keep:
        prs.slides.add_slide(prs.slide_layouts[i])
    master = prs.slide_masters[0]
    id_list = master.slide_layouts._sldLayoutIdLst
    for pos, element in reversed(list(enumerate(list(id_list)))):
        if pos not in keep:
            id_list.remove(element)
    tpl = ROOT / "templates" / "발표자료.pptx"
    prs.save(str(tpl))

    d = ROOT / "docs" / "발표자료"
    d.mkdir(parents=True, exist_ok=True)
    deck(d / "통과.pptx", [0, 1, 1, 5])          # 템플릿이 정의한 것만 씀
    deck(d / "다른템플릿.pptx", [0, 3, 6])        # 템플릿 밖 레이아웃을 씀
    deck(d / "한장만.pptx", [0])                 # 경계: 한 장이지만 템플릿 안쪽


def main() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    build_md()
    build_xlsx()
    build_pptx()
    print(f"픽스처를 만들었다: {ROOT}")


if __name__ == "__main__":
    main()
