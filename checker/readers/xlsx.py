"""엑셀 리더.

시트 이름과, 지정된 자리의 행을 읽는다. 헤더가 몇 행인지는 리더가 추측하지 않는다.
실제 템플릿을 뜯어보니 시트마다 헤더 위치가 다르고, 헤더처럼 보이지만 레이블-값 쌍인
행이 위에 있어서 어떤 heuristic 도 헛짚는다. 그래서 규칙 파일이 시트와 행을 지정한다.
"""
from __future__ import annotations

from pathlib import Path

from checker.readers import reader


@reader(".xlsx", ".xlsm")
def read_xlsx(path: Path) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        return {"sheets": list(wb.sheetnames), "_path": str(path)}
    finally:
        wb.close()


def read_row(path: Path, sheet: str, row: int) -> list[str] | None:
    """한 시트의 한 행을 값 목록으로 읽는다. 시트가 없으면 None.

    꼬리의 빈 칸은 버린다. 엑셀은 쓰다 지운 자리를 빈 셀로 남겨 두는 일이 잦아서,
    그대로 두면 헤더 길이가 파일마다 달라진다.
    """
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            return None
        ws = wb[sheet]
        values = []
        for r in ws.iter_rows(min_row=row, max_row=row, values_only=True):
            values = ["" if v is None else str(v).strip() for v in r]
            break
        while values and values[-1] == "":
            values.pop()
        return values
    finally:
        wb.close()
